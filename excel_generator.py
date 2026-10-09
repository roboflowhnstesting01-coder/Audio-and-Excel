"""Excel creation and formatting."""
from __future__ import annotations

import io
import re
from datetime import datetime
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

import config
from transcription import parse_slide_name

COL_SLIDE = "Slide Name"
COL_NAME = "Audio version"
COL_TEXT = "Audio to text"
COLUMNS = [COL_SLIDE, COL_NAME, COL_TEXT]
MERGED_COLUMNS = ["S.No", COL_SLIDE, COL_NAME, COL_TEXT]
SHEET_NAME = "Transcriptions"
_AUDIO_EXTENSIONS = "mp3|wav|m4a|aac|flac|ogg|wma|aiff|aif|mp4|mpeg|mpga|webm"
_AUDIO_SUFFIX_RE = re.compile(r"\s*-\s*A\s*\d+\s*$", flags=re.IGNORECASE)


@dataclass
class MergeResult:
    """Merged records and review information from uploaded workbooks."""

    dataframe: pd.DataFrame
    files_received: int = 0
    files_processed: int = 0
    worksheets_processed: int = 0
    issues: list[str] = field(default_factory=list)
    duplicates: pd.DataFrame = field(default_factory=pd.DataFrame)

EXCEL_CELL_LIMIT = 32767   # hard limit of characters per Excel cell
_CHUNK_SIZE = 32000        # safe chunk size used when a transcript exceeds the limit


class ExcelGenerationError(Exception):
    """Raised when the Excel file cannot be produced."""


def _is_blank(value) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _header_key(value) -> str:
    return "".join(str(value).split()).casefold()


def _audio_version(value) -> str:
    label = str(value).strip()
    return re.sub(rf"\.(?:{_AUDIO_EXTENSIONS})$", "", label, flags=re.IGNORECASE)


def _audio_slide_name(value) -> str:
    """Return the slide part of a full audio label when it has an A-number suffix."""
    label = _audio_version(value)
    match = _AUDIO_SUFFIX_RE.search(label)
    return label[:match.start()].strip() if match else ""


def _slide_group_key(slide_name, audio_version, row_index: int | None = None) -> str | None:
    """Match by an explicit slide identifier, or by the slide portion of the audio label."""
    name = str(slide_name).strip() if not _is_blank(slide_name) else _audio_slide_name(audio_version)
    if not name:
        return None
    identifier = re.match(r"^(\d+(?:\.\d+)+)(?=\s|$)", name)
    if identifier:
        return f"id:{identifier.group(1).casefold()}"
    return f"name:{_group_key(name)}"


def _worksheet_records(
    dataframe: pd.DataFrame, filename: str, sheet_name: str, issues: list[str]
) -> list[dict[str, str]]:
    if dataframe.empty and len(dataframe.columns) == 0:
        issues.append(f"{filename} / {sheet_name}: worksheet is empty and was skipped.")
        return []

    column_indices = {}
    for index, header in enumerate(dataframe.columns):
        key = _header_key(header)
        if key in {_header_key(name) for name in COLUMNS} and key not in column_indices:
            column_indices[key] = index
    required = {_header_key(name): name for name in COLUMNS}
    missing = [name for key, name in required.items() if key not in column_indices]
    if missing:
        issues.append(
            f"{filename} / {sheet_name}: missing required header(s): {', '.join(missing)}."
        )
        return []
    if dataframe.empty:
        issues.append(f"{filename} / {sheet_name}: worksheet has headers but no records and was skipped.")
        return []

    slide_key = _header_key(COL_SLIDE)
    audio_key = _header_key(COL_NAME)
    text_key = _header_key(COL_TEXT)
    last_slide = ""
    previous_group = None
    records = []
    for position in range(len(dataframe)):
        values = dataframe.iloc[position]
        slide_value = values.iloc[column_indices[slide_key]]
        audio_value = values.iloc[column_indices[audio_key]]
        text_value = values.iloc[column_indices[text_key]]
        if all(_is_blank(value) for value in (slide_value, audio_value, text_value)):
            continue
        if _is_blank(audio_value):
            issues.append(
                f"{filename} / {sheet_name}, row {position + 2}: missing Audio version; row was skipped."
            )
            continue

        if not _is_blank(slide_value):
            last_slide = _clean(str(slide_value).strip())
        else:
            inferred_slide = _clean(_audio_slide_name(audio_value))
            if inferred_slide:
                inferred_key = _slide_group_key(inferred_slide, audio_value)
                previous_key = _slide_group_key(last_slide, audio_value) if last_slide else None
                if inferred_key != previous_key:
                    last_slide = inferred_slide
        effective_slide = last_slide
        group = _slide_group_key(effective_slide, audio_value)
        if group is None:
            issues.append(
                f"{filename} / {sheet_name}, row {position + 2}: slide group could not be determined; review this row."
            )
            display_slide = ""
        else:
            display_slide = effective_slide if group != previous_group else ""
        previous_group = group

        text = "" if _is_blank(text_value) else str(text_value)
        text = re.sub(r"<br\s*/?\s*>", "\n", text, flags=re.IGNORECASE)
        records.append(
            {
                COL_SLIDE: display_slide,
                COL_NAME: _clean(_audio_version(audio_value)),
                COL_TEXT: _clean(text),
                "_effective_slide": effective_slide,
                "_source_file": filename,
                "_worksheet": sheet_name,
                "_source_row": str(position + 2),
            }
        )
    return records


def merge_excel_files(files: Iterable[tuple[str, bytes]]) -> MergeResult:
    """Read valid worksheets from uploaded .xlsx/.xls files in source order."""
    uploaded = list(files)
    result = MergeResult(dataframe=pd.DataFrame(columns=COLUMNS), files_received=len(uploaded))
    records = []
    for filename, content in uploaded:
        try:
            worksheets = pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=object)
        except Exception as exc:  # noqa: BLE001 - report bad uploads in the UI
            result.issues.append(f"{filename}: could not read workbook ({exc}).")
            continue

        result.files_processed += 1
        for sheet_name, dataframe in worksheets.items():
            issue_start = len(result.issues)
            sheet_records = _worksheet_records(dataframe, filename, sheet_name, result.issues)
            new_issues = result.issues[issue_start:]
            missing_headers = any("missing required header(s)" in issue for issue in new_issues)
            empty_without_headers = dataframe.empty and len(dataframe.columns) == 0
            if not missing_headers and not empty_without_headers:
                result.worksheets_processed += 1
            records.extend(sheet_records)

    result.dataframe = pd.DataFrame(records, columns=COLUMNS + [
        "_effective_slide", "_source_file", "_worksheet", "_source_row"
    ])
    if records:
        duplicate_keys = {}
        for index, row in result.dataframe.iterrows():
            key = (_slide_group_key(row["_effective_slide"], row[COL_NAME]), _header_key(row[COL_NAME]))
            duplicate_keys.setdefault(key, []).append(index)
        duplicate_indices = [index for indices in duplicate_keys.values() if len(indices) > 1 for index in indices]
        if duplicate_indices:
            duplicates = result.dataframe.loc[duplicate_indices]
            result.duplicates = duplicates[
                [COL_SLIDE, COL_NAME, COL_TEXT, "_source_file", "_worksheet", "_source_row"]
            ].rename(
                columns={
                    "_source_file": "Source file",
                    "_worksheet": "Worksheet",
                    "_source_row": "Source row",
                }
            ).reset_index(drop=True)
        result.dataframe = result.dataframe[COLUMNS]
    return result


def merged_preview_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Build the four-column display view with one serial number per consecutive slide group."""
    if list(dataframe.columns) != COLUMNS:
        raise ExcelGenerationError("Merged table must contain exactly: " + ", ".join(COLUMNS) + ".")
    rows = []
    serial = 0
    previous_group = None
    for index, values in enumerate(dataframe.itertuples(index=False, name=None)):
        slide_name, audio_version, text = values
        group = _slide_group_key(slide_name, audio_version)
        if group is None and _is_blank(slide_name):
            group = previous_group
        if group is None:
            group = f"unresolved:{index}"
        is_new_group = group != previous_group
        if is_new_group:
            serial += 1
        display_slide = slide_name if not _is_blank(slide_name) else _audio_slide_name(audio_version)
        if not display_slide:
            display_slide = None
        rows.append(
            [serial if is_new_group else None, display_slide if is_new_group else None, audio_version, text]
        )
        previous_group = group
    return pd.DataFrame(rows, columns=MERGED_COLUMNS, dtype=object)


def generate_merged_excel_bytes(dataframe: pd.DataFrame) -> bytes:
    """Create the standardized four-column workbook with export-only serial numbers."""
    if list(dataframe.columns) != COLUMNS:
        raise ExcelGenerationError("Merged table must contain exactly: " + ", ".join(COLUMNS) + ".")
    if dataframe.empty:
        raise ExcelGenerationError("There are no valid audio records to export.")

    try:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = SHEET_NAME
        worksheet.append(MERGED_COLUMNS)
        preview = merged_preview_dataframe(dataframe)
        for row in preview.itertuples(index=False, name=None):
            worksheet.append(
                [row[0], *[None if _is_blank(value) else str(value) for value in row[1:]]]
            )

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        worksheet.column_dimensions["A"].width = 9
        worksheet.column_dimensions["B"].width = 28
        worksheet.column_dimensions["C"].width = 38
        worksheet.column_dimensions["D"].width = 100
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    except Exception as exc:  # noqa: BLE001
        raise ExcelGenerationError(f"Could not create the merged Excel file: {exc}") from exc


def _clean(value: str) -> str:
    """Remove control characters that Excel/openpyxl cannot store."""
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _split_text(text: str, size: int = _CHUNK_SIZE) -> list[str]:
    """Split very long text into <= size chunks, preferring whitespace boundaries. Nothing is dropped."""
    if len(text) <= size:
        return [text]
    chunks, rest = [], text
    while len(rest) > size:
        cut = rest.rfind(" ", int(size * 0.9), size)
        cut = cut if cut != -1 else size
        chunks.append(rest[:cut])
        rest = rest[cut:].lstrip(" ")
    chunks.append(rest)
    return chunks


def _group_key(slide_name: str) -> str:
    """Slides match regardless of capitalization or stray spaces."""
    return " ".join(slide_name.split()).casefold()


def oversized_files(results: Iterable) -> list[str]:
    """Filenames whose transcript is too long for a single Excel cell."""
    return [r.filename for r in results if r.success and len(_clean(r.text)) > EXCEL_CELL_LIMIT]


def build_dataframe(results: Iterable) -> pd.DataFrame:
    """Three-column table of SUCCESSFUL transcriptions only.

        - 'Slide Name' excludes the extension and trailing A1/A2/... suffix; 'Audio version'
            is the full filename stem, including its suffix.
        - Slides remain in first-seen order, and audio files within each slide are ordered
            A1, A2, A3 ...
    - One row per audio file. A transcript longer than Excel's per-cell limit continues on
            following rows with a blank repeated slide name, so no text is lost and no column is added.
    """
    entries = []
    for index, r in enumerate(results):
        if not r.success:
            continue
        slide, number = parse_slide_name(r.filename)
        entries.append(
            {
                "index": index,
                "slide": _clean(slide),
                "audio_version": _clean(r.audio_version),
                "number": number,
                "text": _clean(r.text),
            }
        )

    first_seen: dict[str, int] = {}
    display: dict[str, str] = {}
    for e in entries:
        key = _group_key(e["slide"])
        first_seen.setdefault(key, e["index"])
        display.setdefault(key, e["slide"])  # spelling of the first file seen for that slide
        e["key"] = key

    entries.sort(
        key=lambda e: (
            first_seen[e["key"]],
            e["number"] if e["number"] is not None else 0,
            e["index"],
        )
    )

    rows = []
    written_slides = set()
    for e in entries:
        key = e["key"]
        name = display[key] if key not in written_slides else ""
        for chunk in _split_text(e["text"]):
            rows.append({COL_SLIDE: name, COL_NAME: e["audio_version"], COL_TEXT: chunk})
            name = ""
        written_slides.add(key)
    return pd.DataFrame(rows, columns=COLUMNS)


def generate_excel_bytes(df: pd.DataFrame) -> bytes:
    """Create the formatted .xlsx in memory."""
    if list(df.columns) != COLUMNS:
        raise ExcelGenerationError("Table must have exactly the columns: " + ", ".join(f"'{c}'" for c in COLUMNS) + ".")
    if df.empty:
        raise ExcelGenerationError("There are no successful transcriptions to export.")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(COLUMNS)
        for row in df.itertuples(index=False):
            ws.append([str(v) for v in row])

        # Treat values that start with '=' as plain text, never as formulas.
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        name_width = min(max(max(len(str(v)) for v in df[COL_SLIDE]) + 2, 24), 70)
        ws.column_dimensions["A"].width = name_width
        ws.column_dimensions["B"].width = name_width
        ws.column_dimensions["C"].width = 110
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    except ExcelGenerationError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExcelGenerationError(f"Could not create the Excel file: {exc}") from exc


def save_excel_file(excel_bytes: bytes, output_dir: Optional[Path] = None) -> Path:
    """Write the workbook to output/ with a timestamped name and return its path."""
    folder = Path(output_dir) if output_dir else config.OUTPUT_DIR
    try:
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"transcriptions_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        path.write_bytes(excel_bytes)
        return path
    except OSError as exc:
        raise ExcelGenerationError(f"Could not save the Excel file to '{folder}': {exc}") from exc
